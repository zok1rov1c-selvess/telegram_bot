"""Turli fayl formatlaridan matn ajratib olish."""
from __future__ import annotations

import csv
import io
import logging
import re
from pathlib import Path

from config import MAX_TEXT_CHARS

log = logging.getLogger(__name__)

SUPPORTED = {
    ".pdf", ".docx", ".doc", ".pptx", ".txt", ".md", ".csv", ".json",
    ".xlsx", ".xls", ".html", ".htm", ".rtf", ".log", ".py", ".js",
    ".jpg", ".jpeg", ".png", ".webp", ".bmp",
}

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".bmp"}


# ------------------------------------------------------------------ helpers
def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\x0b\f]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    # ketma-ket takrorlangan qatorlar (kolontitullar) ni siyraklashtirish
    lines, out, prev = text.split("\n"), [], None
    for ln in lines:
        s = ln.strip()
        if s and s == prev:
            continue
        prev = s
        out.append(s)
    return "\n".join(out).strip()


# ------------------------------------------------------------------ parsers
def _from_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages, 1):
        try:
            t = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            t = ""
        if t.strip():
            parts.append(f"\n--- [{i}-bet] ---\n{t}")
    return "\n".join(parts)


def _from_docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        style = (p.style.name or "").lower()
        parts.append(f"## {txt}" if "heading" in style or "title" in style else txt)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _from_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"\n--- [{i}-slayd] ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _from_xlsx(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"\n--- [{ws.title}] ---")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None]
            if vals:
                parts.append(" | ".join(vals))
    wb.close()
    return "\n".join(parts)


def _from_csv(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    try:
        dialect = csv.Sniffer().sniff(raw[:2000])
        rows = list(csv.reader(io.StringIO(raw), dialect))
    except Exception:  # noqa: BLE001
        rows = list(csv.reader(io.StringIO(raw)))
    return "\n".join(" | ".join(r) for r in rows if any(r))


def _from_html(path: Path) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return soup.get_text("\n")


def _from_rtf(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    raw = re.sub(r"\\'([0-9a-fA-F]{2})", "", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", raw)
    return raw.replace("{", " ").replace("}", " ")


def _from_image(path: Path) -> str:
    """OCR (agar tesseract o'rnatilgan bo'lsa)."""
    try:
        import pytesseract  # type: ignore
        from PIL import Image

        return pytesseract.image_to_string(Image.open(path), lang="uzb+rus+eng")
    except Exception as exc:  # noqa: BLE001
        log.info("OCR mavjud emas: %s", exc)
        return ""


def _from_plain(path: Path) -> str:
    for enc in ("utf-8", "cp1251", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


PARSERS = {
    ".pdf": _from_pdf,
    ".docx": _from_docx,
    ".pptx": _from_pptx,
    ".xlsx": _from_xlsx,
    ".xls": _from_xlsx,
    ".csv": _from_csv,
    ".html": _from_html,
    ".htm": _from_html,
    ".rtf": _from_rtf,
}


def extract_text(path: str | Path) -> str:
    """Fayldan toza matn qaytaradi. Blocking - to_thread bilan chaqiring."""
    path = Path(path)
    ext = path.suffix.lower()

    if ext in IMAGE_EXT:
        text = _from_image(path)
    elif ext in PARSERS:
        text = PARSERS[ext](path)
    else:
        text = _from_plain(path)

    text = clean_text(text)
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS] + "\n\n[... matn uzunligi chegarasi ...]"
    return text


def is_supported(filename: str) -> bool:
    return Path(filename).suffix.lower() in SUPPORTED
